"""Ekstraksi teks dan OCR visual dari file (PDF, PPTX, DOCX, Gambar)."""
import base64
import io
from typing import Callable, List, Tuple

import fitz  # PyMuPDF
from docx import Document
from langchain_core.messages import HumanMessage
from pptx import Presentation

from core.ai import to_text, vision_llm


def read_image(file_bytes: bytes, file_extension: str) -> List[Tuple[str, str]]:
    # Transkripsi gambar/diagram/rumus menggunakan Vision LLM
    mime_type = "jpeg" if file_extension.lower() in ("jpg", "jpeg") else file_extension.lower()
    base64_image = base64.b64encode(file_bytes).decode("utf-8")

    prompt_message = HumanMessage(content=[
        {
            "type": "text",
            "text": "Transkripsikan semua teks, judul, poin-poin, diagram, bagan, dan rumus pada gambar ini secara lengkap dalam Bahasa Indonesia untuk materi kuliah RAG.",
        },
        {
            "type": "image_url",
            "image_url": f"data:image/{mime_type};base64,{base64_image}",
        },
    ])

    transcription = to_text(vision_llm.invoke([prompt_message]))
    return [(transcription, "Gambar")]


def read_pdf(file_bytes: bytes, file_extension: str) -> List[Tuple[str, str]]:
    # Baca PDF per halaman; jika teks minim (< 100 char), gunakan Vision OCR
    pdf_document = fitz.open(stream=file_bytes, filetype="pdf")
    results: List[Tuple[str, str]] = []

    for page_num, page in enumerate(pdf_document, 1):
        page_text = (page.get_text() or "").strip()

        # OCR visual untuk halaman diagram/infografis
        if len(page_text) < 100:
            try:
                pixmap = page.get_pixmap(dpi=130)
                ocr_text, _ = read_image(pixmap.tobytes("png"), "png")[0]
                if page_text and ocr_text:
                    page_text = f"{page_text}\n\n[Transkripsi Visual & Infografis]:\n{ocr_text}"
                elif ocr_text:
                    page_text = ocr_text
            except Exception as error:
                print(f"[Warning] OCR gagal pada PDF Hal {page_num}: {error}")

        if page_text:
            results.append((page_text, f"Hal {page_num}"))

    pdf_document.close()
    return results


def read_pptx(file_bytes: bytes, file_extension: str) -> List[Tuple[str, str]]:
    # Baca PowerPoint slide per slide (judul, teks, tabel, gambar tertanam, catatan)
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(file_bytes))
    results: List[Tuple[str, str]] = []

    for slide_num, slide in enumerate(presentation.slides, 1):
        slide_parts = []
        image_counter = 0

        # Judul slide
        if slide.shapes.title and slide.shapes.title.text:
            slide_parts.append(f"Judul: {slide.shapes.title.text.strip()}")

        # Kotak teks, tabel, dan gambar tertanam
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue

            if shape.has_text_frame and shape.text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_parts.append(text)

            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        slide_parts.append(row_text)

            # Ekstrak gambar tertanam pada slide dan jalankan OCR
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img_blob = shape.image.blob
                    img_ext = shape.image.content_type.split("/")[-1]
                    image_counter += 1
                    ocr_text, _ = read_image(img_blob, img_ext)[0]
                    if ocr_text.strip():
                        slide_parts.append(f"[Gambar {image_counter} - Transkripsi Visual]:\n{ocr_text.strip()}")
                except Exception as img_error:
                    print(f"[Warning] Gagal ekstrak gambar slide {slide_num}: {img_error}")

        # Catatan presenter jika ada
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"[Catatan: {notes}]")

        full_text = "\n".join(slide_parts).strip()
        if full_text:
            results.append((full_text, f"Slide {slide_num}"))

    return results


def read_docx(file_bytes: bytes, file_extension: str) -> List[Tuple[str, str]]:
    # Baca dokumen Word per 12 paragraf, tabel, dan gambar tertanam
    docx_document = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in docx_document.paragraphs if p.text.strip()]
    results: List[Tuple[str, str]] = []

    # Kelompokkan paragraf per 12 bagian
    batch_size = 12
    for start in range(0, len(paragraphs), batch_size):
        sec_num = start // batch_size + 1
        sec_text = "\n".join(paragraphs[start:start + batch_size]).strip()
        if sec_text:
            results.append((sec_text, f"Bagian {sec_num}"))

    # Ekstrak isi tabel
    for table_num, table in enumerate(docx_document.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells if c.text.strip()) for row in table.rows]
        table_text = "\n".join(r for r in rows if r)
        if table_text:
            results.append((table_text, f"Tabel {table_num}"))

    # Ekstrak gambar tertanam dari relasi dokumen
    image_counter = 0
    try:
        for rel in docx_document.part.rels.values():
            if "image" in rel.reltype:
                try:
                    img_part = rel.target_part
                    img_ext = img_part.content_type.split("/")[-1] if "/" in img_part.content_type else "png"
                    image_counter += 1
                    ocr_text, _ = read_image(img_part.blob, img_ext)[0]
                    if ocr_text.strip():
                        results.append((f"[Gambar {image_counter} - Transkripsi Visual]:\n{ocr_text.strip()}", f"Gambar {image_counter}"))
                except Exception as img_error:
                    print(f"[Warning] Gagal ekstrak gambar DOCX #{image_counter}: {img_error}")
    except Exception as rel_error:
        print(f"[Warning] Gagal akses relasi gambar DOCX: {rel_error}")

    return results


# Peta ekstensi berkas ke fungsi pembacanya
READERS: dict[str, Callable] = {
    "pdf": read_pdf,
    "pptx": read_pptx,
    "docx": read_docx,
    "png": read_image,
    "jpg": read_image,
    "jpeg": read_image,
}


def load_file(filename: str, file_bytes: bytes) -> Tuple[List[Tuple[str, str]], str]:
    # Deteksi ekstensi file dan ekstrak teks beserta label halamannya
    file_extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if file_extension not in READERS:
        supported = ", ".join(f".{ext}" for ext in READERS)
        raise ValueError(f"Format .{file_extension} tidak didukung. Format yang didukung: {supported}")

    return READERS[file_extension](file_bytes, file_extension), file_extension

